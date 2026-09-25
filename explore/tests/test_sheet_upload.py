"""#177: a spreadsheet of items → HWDB. The sheet reader, mapping, plan and
per-row apply are exercised on their own; the view with HWDB mocked.

    python manage.py test explore.tests.test_sheet_upload
"""

from __future__ import annotations

import io
from unittest import mock

import openpyxl
import requests
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings

from explore import sheetupload as su
from explore.models import ActivityEvent, HwdbComponentEvent, SheetJob
from hwdb.fnal.bearer import FnalLinkRequired

T = "D00400300001"
C = "D00400300002"
URL = f"/hw/dev/upload/{T}/"
TEMPLATE = {"Vbd": None, "Notes": None, "DATA": {}}
CONNECTORS = {"A1": C, "A2": C}
INSTITUTIONS = [{"id": 186, "name": "BNL", "country_code": "US"}, {"id": 7, "name": "CERN", "country_code": "CH"}]
MAKERS = [{"id": 7, "name": "HPK"}, {"id": 8, "name": "SMB"}]

CSV = (b"Record Type,Item\r\nPart Type ID,D00400300001\r\nInstitution,(186) BNL\r\n\r\n"
       b"External ID,Serial Number,Status,S:Vbd,C:A1,Comments,Extra\r\n"
       b",HPK-9,110,51.5,,new one,x\r\n"
       b"D00400300001-00001,HPK-1,QA/QC Tests - Passed All,52,D00400300002-00001,,\r\n"
       b",,,,,,\r\n"
       b"D00400300001-00001,,,,, second row,\r\n")


def _live(n, sn, status=0, vbd=51.4, comments=""):
    return {"part_id": f"{T}-{n:05d}", "serial_number": sn, "status": {"id": status, "name": "x"},
            "comments": comments, "specifications": [{"Vbd": vbd, "Notes": None, "DATA": {}}]}


LIVE = {f"{T}-00001": _live(1, "HPK-1"), f"{T}-00002": _live(2, "HPK-2", 120),
        f"{T}-00004": _live(4, "DUP"), f"{T}-00005": _live(5, "DUP")}


class ReadTest(TestCase):
    def test_csv_with_the_utilitys_header_block(self):
        (s,) = su.read_sheets("items.csv", CSV)
        self.assertEqual(s["values"], {"Record Type": "Item", "Part Type ID": T, "Institution": "(186) BNL"})
        self.assertEqual(s["columns"], ["External ID", "Serial Number", "Status", "S:Vbd", "C:A1", "Comments", "Extra"])
        self.assertEqual([r[0] for r in s["rows"]], [6, 7, 9])          # sheet row numbers, the blank row dropped
        self.assertEqual(s["rows"][0][1], [None, "HPK-9", 110, 51.5, None, "new one", "x"])
        self.assertEqual(s["rows"][1][1][3], 52)

    def test_plain_table_and_typed_cells(self):
        (s,) = su.read_sheets("a.csv", b"Serial Number;S:Vbd\n00123;1e3\nabc;0.5\n")
        self.assertEqual(s["values"], {})
        self.assertEqual(s["rows"], [[2, ["00123", 1000.0]], [3, ["abc", 0.5]]])   # a leading zero stays text
        self.assertEqual(su._cell(" <null> "), "<null>")
        self.assertEqual(su._cell("-7"), -7)
        self.assertEqual(su._cell(3.0), 3)
        self.assertIsNone(su._cell(""))

    def test_workbook_tabs(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Items"
        ws.append(["Serial Number", "Vbd"])
        ws.append(["HPK-1", 51.0])
        ws.append(["HPK-2", 52.25])
        wb.create_sheet("Empty")
        buf = io.BytesIO()
        wb.save(buf)
        sheets = su.read_sheets("book.xlsx", buf.getvalue())
        self.assertEqual([s["name"] for s in sheets], ["Items"])
        self.assertEqual(sheets[0]["rows"], [[2, ["HPK-1", 51]], [3, ["HPK-2", 52.25]]])

    def test_json_in_a_cell_and_json_files(self):
        (s,) = su.read_sheets("a.csv", b'Serial Number,T:Curve,T:Meta\nHPK-1,"[1, 2, 3]","{""a"": {""b"": 1}}"\nHPK-2,{oops,x\n')
        self.assertEqual(s["rows"][0][1], ["HPK-1", [1, 2, 3], {"a": {"b": 1}}])
        self.assertEqual(s["rows"][1][1][1], "{oops")                  # not JSON: stays text
        (s,) = su.read_sheets("r.json", b'[{"Serial Number": "HPK-1", "Gain": 12, "Curve": [1, 2]}, {"Serial Number": "HPK-2", "Noise": {"rms": 0.5}}]')
        self.assertEqual((s["name"], s["values"], s["columns"]), ("r", {}, ["Serial Number", "Gain", "Curve", "Noise"]))
        self.assertEqual(s["rows"], [[1, ["HPK-1", 12, [1, 2], None]], [2, ["HPK-2", None, None, {"rms": 0.5}]]])
        (s,) = su.read_sheets("k.json", b'{"D00400300001-00001": {"Gain": 1}, "D00400300001-00002": {"Gain": 2}, "note": "x"}')
        self.assertEqual(s["columns"], ["External ID", "Gain"])
        self.assertEqual(s["rows"][1][1], [f"{T}-00002", 2])
        (s,) = su.read_sheets("w.json", b'{"Record Type": "Test", "Test Name": "QC", "data": [{"External ID": "X", "v": 1}]}')
        self.assertEqual((s["values"], s["columns"], s["rows"]), ({"Record Type": "Test", "Test Name": "QC"}, ["External ID", "v"], [[1, ["X", 1]]]))
        self.assertEqual(su.read_sheets("e.json", b'[]'), [])

    def test_a_block_is_two_columns_wide(self):
        (s,) = su.read_sheets("a.csv", b"Serial Number,Vbd,Notes\nHPK-1,1,\n\nHPK-2,2,\n")
        self.assertEqual(s["values"], {})                       # three columns: a table with a gap, not a block
        self.assertEqual(len(s["rows"]), 2)
        self.assertEqual(su.read_sheets("a.csv", b"\n\n"), [])


class MapTest(TestCase):
    def test_auto_map(self):
        m = su.auto_map(["External ID", "serial number", "S:Vbd", "c:A1", "Vbd", "a2", "Extra", "Part Type ID"],
                        TEMPLATE, CONNECTORS)
        self.assertEqual(m, {"External ID": "part_id", "serial number": "serial_number", "S:Vbd": "spec:Vbd",
                             "c:A1": "pos:A1", "Vbd": "spec:Vbd", "a2": "pos:A2", "Extra": "",
                             "Part Type ID": "part_type_id"})

    def test_records_merge_rows_and_take_the_sheet_defaults(self):
        (s,) = su.read_sheets("items.csv", CSV)
        recs = su.records(s, su.auto_map(s["columns"], TEMPLATE, CONNECTORS))
        self.assertEqual(len(recs), 2)
        new, old = recs
        self.assertEqual((new["part_id"], new["serial_number"], new["institution"], new["part_type_id"]),
                         ("", "HPK-9", "(186) BNL", T))
        self.assertEqual((new["specs"], new["positions"], new["comments"]), ({"Vbd": 51.5}, {}, "new one"))
        self.assertEqual(old["rows"], [7, 9])
        self.assertEqual(old["comments"], "second row")                  # a later non-blank cell wins
        self.assertEqual(old["positions"], {"A1": f"{C}-00001"})
        self.assertNotIn("Extra", old)                                    # unassigned column ignored

    def test_test_columns_and_kind(self):
        m = su.auto_map(["Serial Number", "T:Gain", "gain", "Noise", "Comments", "Other"], TEMPLATE, CONNECTORS, ["Gain", "Noise"])
        self.assertEqual(m, {"Serial Number": "serial_number", "T:Gain": "test:Gain", "gain": "test:Gain",
                             "Noise": "test:Noise", "Comments": "comments", "Other": ""})
        self.assertEqual(su.auto_map(["Other", "Vbd"], TEMPLATE, CONNECTORS, (), "test"), {"Other": "test:Other", "Vbd": "spec:Vbd"})
        self.assertEqual(su.detect_kind({"values": {"Record Type": "Test", "Test Name": "QC"}, "columns": ["Serial Number"]}),
                         ("test", "QC"))
        self.assertEqual(su.detect_kind({"values": {}, "columns": ["Serial Number", "T:x"]}), ("test", ""))
        self.assertEqual(su.detect_kind({"values": {"Record Type": "Item"}, "columns": ["T:x"]}), ("test", ""))
        self.assertEqual(su.detect_kind({"values": {}, "columns": ["Serial Number", "S:x"]}), ("item", ""))
        self.assertEqual(su.detect_kind({"values": {"Record Type": "Item Image"}, "columns": ["Serial Number"]}), ("image", ""))
        self.assertEqual(su.detect_kind({"values": {"Record Type": "Test Image", "Test Name": "QC"}, "columns": ["x"]}), ("image", "QC"))
        self.assertEqual(su.detect_kind({"values": {}, "columns": ["External ID", "Image File", "Comments"]}), ("image", ""))

    def test_records_key_on_pid_and_serial(self):
        s = {"values": {}, "columns": ["Part ID", "Serial Number"],
             "rows": [[2, [f"{T}-00001", None]], [3, [f"{T}-00001", None]], [4, ["<unassigned>", "X"]], [5, [None, None]]]}
        recs = su.records(s, su.auto_map(s["columns"], {}, {}))
        self.assertEqual([(r["part_id"], r["serial_number"], r["rows"]) for r in recs],
                         [(f"{T}-00001", "", [2, 3]), ("", "X", [4]), ("", "", [5])])


def _plan(rows, **kw):
    s = {"values": kw.pop("values", {}), "columns": kw.pop("columns"), "rows": [[i + 2, r] for i, r in enumerate(rows)]}
    recs = su.records(s, su.auto_map(s["columns"], TEMPLATE, CONNECTORS))
    return su.plan(recs, T, kw.pop("live", LIVE), kw.pop("makers", {f"{T}-00001": "HPK", f"{T}-00002": "HPK"}), TEMPLATE, CONNECTORS,
                   INSTITUTIONS, MAKERS)


class PlanTest(TestCase):
    COLS = ["External ID", "Serial Number", "Status", "Manufacturer", "S:Vbd", "S:DATA.received", "C:A1",
            "Institution", "Location", "Comments"]

    def test_create_patch_skip(self):
        rows = _plan([
            [None, "HPK-9", 110, "HPK", 51.5, None, None, "(186) BNL", None, "new"],   # unknown serial → create
            [f"{T}-00001", None, "QA/QC Tests - Passed All", "SMB", 52, 3, f"{C}-00001", None, "CERN", None],
            [None, "hpk-2", 120, "hpk", 51.4, None, None, None, None, ""],             # everything as it is
        ], columns=self.COLS)
        self.assertEqual([(r["action"], r["state"]) for r in rows],
                         [("create", "pending"), ("patch", "pending"), ("skip", "done")])
        new, upd, same = rows
        self.assertEqual(new["changes"], ["new item"])
        self.assertEqual((new["rec"]["institution"]["id"], new["rec"]["status_id"], new["rec"]["manufacturer"]["id"]),
                         (186, 110, 7))
        self.assertEqual(upd["pid"], f"{T}-00001")
        self.assertEqual(upd["changes"], ["status → QA/QC Tests - Passed All", "manufacturer → SMB", "Vbd → 52",
                                          "DATA.received → 3", "location → CERN", f"A1 → {C}-00001"])
        self.assertEqual(same["pid"], f"{T}-00002")

    def test_errors(self):
        rows = _plan([
            [f"{T}-00099", None, None, None, None, None, None, None, None, None],
            [f"{C}-00001", None, None, None, None, None, None, None, None, None],
            [None, "DUP", None, None, None, None, None, None, None, None],
            [None, "HPK-1", "Lost", None, None, None, None, None, None, None],
            [f"{T}-00001", None, None, "Acme", None, None, None, None, None, None],
            [f"{T}-00002", None, None, None, None, None, None, None, "Mars", None],
            [None, "HPK-9", None, None, None, None, None, None, None, None],
            [None, None, None, None, 1, None, None, None, None, None],
        ], columns=self.COLS)
        self.assertTrue(all(r["action"] == "error" for r in rows))
        errs = [r["error"] for r in rows]
        self.assertIn("not in HWDB", errs[0])
        self.assertIn(f"is not a {T} item", errs[1])
        self.assertIn("DUP is on 2 items", errs[2])
        self.assertIn("unknown status", errs[3])
        self.assertIn("not a manufacturer", errs[4])
        self.assertIn("unknown location", errs[5])
        self.assertIn("needs an institution", errs[6])
        self.assertIn("no External ID and no serial", errs[7])

    def test_spec_keys_positions_and_type_are_checked(self):
        rows = _plan([[None, "HPK-1", 1]], columns=["External ID", "Serial Number", "S:Gain"])
        self.assertIn("not a key of the type", rows[0]["error"])
        rows = _plan([[None, "HPK-1", 1]], columns=["External ID", "Serial Number", "S:Vbd.x"])
        self.assertIn("not a nested object", rows[0]["error"])
        rows = _plan([[None, "HPK-1", 1]], columns=["External ID", "Serial Number", "C:B9"])
        self.assertIn("no “B9” position", rows[0]["error"])
        rows = _plan([[None, "HPK-1"]], columns=["External ID", "Serial Number"], values={"Part Type ID": C})
        self.assertIn(f"is not {T}", rows[0]["error"])

    def test_sheet_defaults_apply_to_every_row(self):
        rows = _plan([[None, "HPK-9"], [None, "HPK-8"]], columns=["External ID", "Serial Number"],
                     values={"Institution": "186", "Status": "In Fabrication"})
        self.assertEqual([(r["action"], r["rec"]["institution"]["name"], r["rec"]["status_id"]) for r in rows],
                         [("create", "BNL", 100)] * 2)


class PlanTestsTest(TestCase):
    COLS = ["External ID", "Serial Number", "T:Gain", "T:Noise.rms", "Comments"]

    def rows(self, rows, name="QC", **kw):
        s = {"values": kw.get("values", {}), "columns": self.COLS, "rows": [[i + 2, r] for i, r in enumerate(rows)]}
        return su.plan_tests(su.records(s, su.auto_map(self.COLS, {}, {}), merge=False), T, LIVE, name)

    def test_one_record_per_row(self):
        rows = self.rows([
            [None, "HPK-1", 12, 0.5, "first"],
            [None, "HPK-1", 13, None, None],          # a second record on the same item, not merged
            [f"{T}-00002", None, 1, None, None],
        ])
        self.assertEqual([(r["pid"], r["action"], r["state"]) for r in rows],
                         [(f"{T}-00001", "test", "pending")] * 2 + [(f"{T}-00002", "test", "pending")])
        self.assertEqual(rows[0]["rec"], {"test_name": "QC", "comments": "first", "data": {"Gain": 12, "Noise": {"rms": 0.5}}})
        self.assertEqual(rows[1]["rec"]["data"], {"Gain": 13})
        self.assertEqual(rows[0]["changes"], ["QC: Gain = 12, Noise.rms = 0.5"])
        rows = self.rows([[None, "HPK-1", {"a": [1, 2]}, None, None]])
        self.assertEqual(rows[0]["rec"]["data"], {"Gain": {"a": [1, 2]}})
        self.assertEqual(rows[0]["changes"], ['QC: Gain = {"a": [1, 2]}'])

    def test_errors(self):
        rows = self.rows([
            [None, "HPK-9", 1, None, None],
            [None, "DUP", 1, None, None],
            [f"{T}-00099", None, 1, None, None],
            [None, "HPK-1", None, None, "no values"],
            [None, None, 1, None, None],
        ])
        self.assertTrue(all(r["action"] == "error" for r in rows))
        errs = [r["error"] for r in rows]
        self.assertIn("no D00400300001 item has serial number HPK-9", errs[0])
        self.assertIn("DUP is on 2 items", errs[1])
        self.assertIn("not in HWDB", errs[2])
        self.assertIn("no test values", errs[3])
        self.assertIn("no External ID and no serial", errs[4])
        self.assertIn("no test name", self.rows([[None, "HPK-1", 1, None, None]], name="")[0]["error"])
        self.assertEqual(self.rows([[None, "HPK-1", 1, None, None]], name="", values={"Test Name": "From block"})[0]["rec"]["test_name"],
                         "From block")


class GroupTest(TestCase):
    """#179: [] keys build lists — the utility's nested groups."""
    COLS = ["Serial Number", "T:Test Date", "T:Operator", "T:Subtests[].Name", "T:Subtests[].Trials[].Number",
            "T:Subtests[].Trials[].Value", "T:Tags[]"]

    def plan(self, rows):
        s = {"values": {}, "columns": self.COLS, "rows": [[i + 2, r] for i, r in enumerate(rows)]}
        return su.plan_tests(su.records(s, su.auto_map(self.COLS, {}, {}), merge=False), T, LIVE, "QC")

    def test_rows_nest_by_their_scalar_members(self):
        rows = self.plan([
            ["HPK-1", "2026-09-01", "cz", "A", 1, 0.1, "x"],
            ["HPK-1", "2026-09-01", "cz", "A", 2, 0.2, "y"],
            ["HPK-1", "2026-09-01", "cz", "B", 1, 0.3, None],
            ["HPK-1", "2026-09-02", "cz", "A", 1, 0.4, None],     # another date: another record
            ["HPK-2", "2026-09-01", "cz", "A", 1, 0.5, None],     # another item
        ])
        self.assertEqual([(r["pid"], r["rows"]) for r in rows],
                         [(f"{T}-00001", [2, 3, 4]), (f"{T}-00001", [5]), (f"{T}-00002", [6])])
        self.assertEqual(rows[0]["rec"]["data"], {
            "Test Date": "2026-09-01", "Operator": "cz",
            "Subtests": [{"Name": "A", "Trials": [{"Number": 1, "Value": 0.1}, {"Number": 2, "Value": 0.2}]},
                         {"Name": "B", "Trials": [{"Number": 1, "Value": 0.3}]}],
            "Tags": ["x", "y"]})
        self.assertEqual(rows[1]["rec"]["data"]["Subtests"], [{"Name": "A", "Trials": [{"Number": 1, "Value": 0.4}]}])
        self.assertIn('"Subtests": [{"Name": "A"', rows[0]["changes"][0])

    def test_without_list_keys_every_row_is_a_record(self):
        cols = ["Serial Number", "T:a.b"]
        s = {"values": {}, "columns": cols, "rows": [[2, ["HPK-1", 1]], [3, ["HPK-1", 2]]]}
        rows = su.plan_tests(su.records(s, su.auto_map(cols, {}, {}), merge=False), T, LIVE, "QC")
        self.assertEqual([r["rec"]["data"] for r in rows], [{"a": {"b": 1}}, {"a": {"b": 2}}])


class ImageTest(TestCase):
    COLS = ["External ID", "Serial Number", "Image File", "Save As", "Comments", "History Order"]

    def plan(self, rows, name=""):
        s = {"values": {}, "columns": self.COLS, "rows": [[i + 2, r] for i, r in enumerate(rows)]}
        return su.plan_images(su.records(s, su.auto_map(self.COLS, {}, {}), merge=False), T, LIVE, name)

    def test_plan(self):
        rows = self.plan([
            [None, "HPK-1", "photos/front.png", None, "front", None],
            [f"{T}-00002", None, "C:\\scans\\sheet.pdf", "datasheet.pdf", None, 1],
            [None, "HPK-1", None, None, None, None],
            [None, "HPK-9", "x.png", None, None, None],
            [None, "HPK-1", "x.png", None, None, "two"],
        ], name="QC")
        self.assertEqual([(r["pid"], r["action"]) for r in rows[:2]], [(f"{T}-00001", "image"), (f"{T}-00002", "image")])
        self.assertEqual(rows[0]["rec"], {"file": "front.png", "save_as": "front.png", "comments": "front", "test_name": "QC", "hist_order": 0})
        self.assertEqual(rows[0]["changes"], ["front.png → test “QC”"])
        self.assertEqual((rows[1]["rec"]["file"], rows[1]["rec"]["save_as"], rows[1]["rec"]["hist_order"]), ("sheet.pdf", "datasheet.pdf", 1))
        self.assertEqual(rows[1]["changes"], ["sheet.pdf as datasheet.pdf → test “QC” #1"])
        self.assertEqual([r["error"] for r in rows[2:]], ["no image file", f"no {T} item has serial number HPK-9",
                                                          "history order “two” is not a number"])
        self.assertEqual(self.plan([[None, "HPK-1", "a.png", None, None, None]])[0]["changes"], ["a.png → item"])

    def test_apply_item_and_test_attachments(self):
        api = mock.MagicMock()
        api.get_images.return_value = {"data": [{"image_name": "old.png"}]}
        api.post_component_image.return_value = _ok()
        row = {"pid": f"{T}-00001", "rec": {"file": "a.png", "save_as": "a.png", "comments": "", "test_name": "", "hist_order": 0}}
        self.assertEqual(su.apply_image_row(api, row, b"x", "image/png", lambda n: None), (f"{T}-00001", ["attached"]))
        self.assertEqual(api.post_component_image.call_args.args, (f"{T}-00001", b"x", "a.png", "Sheet upload via HWDB Explorer", "image/png"))
        row["rec"]["save_as"] = "old.png"
        self.assertEqual(su.apply_image_row(api, row, b"x", "image/png", lambda n: None)[1], ["already attached"])
        api.get_tests.return_value = {"data": [{"id": 71}, {"id": 70}]}
        api.get_test_images.return_value = {"data": []}
        api.post_test_image.return_value = _ok()
        row["rec"].update(test_name="QC", hist_order=1, save_as="a.png", comments="c")
        self.assertEqual(su.apply_image_row(api, row, b"x", "image/png", lambda n: 5)[1], ["attached"])
        self.assertEqual(api.post_test_image.call_args.args, (70, b"x", "a.png", "c", "image/png"))
        row["rec"]["hist_order"] = 2
        with self.assertRaises(su.SheetError):
            su.apply_image_row(api, row, b"x", "image/png", lambda n: 5)
        with self.assertRaises(su.SheetError):
            su.apply_image_row(api, row, b"x", "image/png", lambda n: None)


class ApplyTestsTest(TestCase):
    def setUp(self):
        self.api = mock.MagicMock()
        self.api.post_test.return_value = _ok()
        self.api.get_tests.return_value = {"data": [{"test_data": {"DATA": {"Gain": 12}}}]}
        self.row = {"pid": f"{T}-00001", "rec": {"test_name": "QC", "comments": "c", "data": {"Gain": 12}}}

    def test_same_record_is_left_alone(self):
        self.assertEqual(su.apply_test_row(self.api, T, self.row, lambda n: 5), (f"{T}-00001", ["already recorded"]))
        self.api.get_tests.assert_called_once_with(f"{T}-00001", test_type_id=5, history=True)
        self.api.post_test.assert_not_called()

    def test_post(self):
        self.row["rec"]["data"] = {"Gain": 13}
        self.assertEqual(su.apply_test_row(self.api, T, self.row, lambda n: 5)[1], ["test posted"])
        self.assertEqual(self.api.post_test.call_args.args,
                         (f"{T}-00001", {"comments": "c", "test_type": "QC", "test_data": {"DATA": {"Gain": 13}}}))
        self.api.get_tests.reset_mock()
        self.assertEqual(su.apply_test_row(self.api, T, self.row, lambda n: None)[1], ["test posted"])   # no id known: no check
        self.api.get_tests.assert_not_called()
        self.api.post_test.return_value = {"status": "ERROR", "data": "roles"}
        with self.assertRaises(su.SheetError):
            su.apply_test_row(self.api, T, self.row, lambda n: 5)


def _detail(n, sn="HPK-1", status=0, vbd=51.4, maker=7, location=186):
    return {"data": {"part_id": f"{T}-{n:05d}", "serial_number": sn, "status": {"id": status},
                     "manufacturer": {"id": maker, "name": "HPK"} if maker else None, "comments": "",
                     "specifications": [{"Vbd": vbd, "Notes": None, "DATA": {}}],
                     "location": {"id": location, "name": "BNL"}}}


def _ok(**kw):
    return {"status": "OK", **kw}


class ApplyTest(TestCase):
    def setUp(self):
        self.api = mock.MagicMock()
        self.api.create_component.return_value = _ok(part_id=f"{T}-00009")
        self.api.patch_component.return_value = _ok()
        self.api.post_location.return_value = _ok()
        self.api.patch_subcomponents.return_value = _ok()
        self.api.get_component.return_value = _detail(1)
        self.api.get_subcomponents.return_value = {"data": [{"functional_position": "A1", "part_id": f"{C}-00001"}]}
        self.lookups = {}

    def lookup(self, tid, sn):
        return self.lookups.get((tid, sn), [])

    def go(self, row):
        return su.apply_row(self.api, T, row, TEMPLATE, CONNECTORS, MAKERS, self.lookup, "2026-09-25T10:00:00")

    def test_create_then_positions(self):
        row = {"action": "create", "pid": "", "rec": {
            "serial_number": "HPK-9", "institution": {"id": 186, "country_code": "US"}, "comments": "new",
            "specs": {"Vbd": 51.5, "DATA.received": 3}, "positions": {"A1": "S-1", "A2": "<null>"}, "status_id": 110}}
        self.lookups[(C, "S-1")] = [f"{C}-00007"]
        self.api.get_subcomponents.return_value = {"data": []}
        pid, done = self.go(row)
        self.assertEqual((pid, done), (f"{T}-00009", ["created", "positions"]))
        payload = self.api.create_component.call_args.args[1]
        self.assertEqual(payload["specifications"], {"Vbd": 51.5, "Notes": None, "DATA": {"received": 3}})
        self.assertEqual((payload["status"], payload["institution"], payload["country_code"], payload["comments"]),
                         ({"id": 110}, {"id": 186}, "US", "new"))
        self.assertNotIn("manufacturer", payload)                  # two makers on the type, none in the row
        self.api.get_component.assert_not_called()
        self.api.patch_component.assert_not_called()
        self.assertEqual(self.api.patch_subcomponents.call_args.args[1]["subcomponents"],
                         {"A1": f"{C}-00007", "A2": None})

    def test_create_finds_the_item_a_lost_reply_created(self):
        row = {"action": "create", "pid": "", "rec": {"serial_number": "HPK-1", "institution": {"id": 186, "country_code": "US"},
                                                      "specs": {}, "positions": {}, "comments": ""}}
        self.lookups[(T, "HPK-1")] = [f"{T}-00001"]
        pid, done = self.go(row)
        self.assertEqual((pid, done), (f"{T}-00001", []))          # nothing differed, nothing written
        self.api.create_component.assert_not_called()
        self.api.patch_component.assert_not_called()
        self.lookups[(T, "HPK-1")] = [f"{T}-00001", f"{T}-00004"]
        with self.assertRaises(su.SheetError):
            self.go(row)

    def test_patch_only_what_differs(self):
        row = {"action": "patch", "pid": f"{T}-00001", "rec": {"serial_number": "", "status_id": 120,
                                                                "specs": {"Vbd": 52}, "positions": {}}}
        pid, done = self.go(row)
        self.assertEqual(done, ["updated"])
        p = self.api.patch_component.call_args.args[1]
        self.assertEqual(p, {"part_id": f"{T}-00001", "serial_number": "HPK-1", "manufacturer": {"id": 7},
                             "specifications": {"Vbd": 52, "Notes": None, "DATA": {}}, "status": {"id": 120},
                             "comments": ""})
        self.api.patch_component.reset_mock()
        row["rec"] = {"serial_number": "HPK-1", "status_id": 0, "specs": {"Vbd": 51.4}, "positions": {},
                      "location": {"id": 186}}
        self.assertEqual(self.go(row)[1], [])                       # the location matches too
        self.api.patch_component.assert_not_called()
        self.api.post_location.assert_not_called()

    def test_location_and_positions_by_serial(self):
        row = {"action": "patch", "pid": f"{T}-00001", "rec": {
            "serial_number": "", "specs": {}, "positions": {"A2": "S-2"}, "location": {"id": 7},
            "location_comments": "moved", "arrived": "2026-09-01T00:00:00"}}
        self.lookups[(C, "S-2")] = [f"{C}-00002"]
        self.assertEqual(self.go(row)[1], ["location", "positions"])
        self.assertEqual(self.api.post_location.call_args.args[1],
                         {"location": {"id": 7}, "arrived": "2026-09-01T00:00:00", "comments": "moved"})
        self.assertEqual(self.api.patch_subcomponents.call_args.args[1]["subcomponents"],
                         {"A1": f"{C}-00001", "A2": f"{C}-00002"})
        self.lookups[(C, "S-2")] = []
        with self.assertRaises(su.SheetError):
            self.go(row)

    def test_refusals_raise(self):
        self.api.patch_component.return_value = {"status": "ERROR", "data": "nope"}
        row = {"action": "patch", "pid": f"{T}-00001", "rec": {"serial_number": "", "status_id": 120, "specs": {}, "positions": {}}}
        with self.assertRaises(su.SheetError) as cm:
            self.go(row)
        self.assertEqual(str(cm.exception), "nope")


def _api():
    api = mock.MagicMock()
    api.get_component_type.return_value = {"data": {
        "manufacturers": MAKERS, "connectors": CONNECTORS,
        "properties": {"specifications": [{"datasheet": TEMPLATE}]}}}
    api.get_institutions.return_value = {"data": [{"id": 186, "name": "BNL", "country": {"code": "US", "name": "USA"}}]}
    api._make_request.return_value = {"pagination": {"pages": 1}, "data": list(LIVE.values())}
    api.create_component.return_value = _ok(part_id=f"{T}-00009")
    api.patch_component.return_value = _ok()
    api.get_component.return_value = _detail(1)
    return api


def _mocked(api):
    return (mock.patch("explore.views.mint_for", return_value="bearer"),
            mock.patch("explore.views.FnalDbApiClient", return_value=api))


class ViewTest(TestCase):
    def setUp(self):
        HwdbComponentEvent.objects.create(instance="dev", part_type_id=T, part_id=f"{T}-00001",
                                          serial_number="HPK-1", status="Unknown", status_id=0, manufacturer="HPK")
        self.client.force_login(get_user_model().objects.create_user("w", "w@w.io", "pw"))

    def upload(self, api, name="items.csv", body=CSV, **extra):
        p1, p2 = _mocked(api)
        with p1, p2:
            return self.client.post(URL, {"step": "file", "file": _file(name, body), **extra})

    def test_page_and_unlinked(self):
        p1, p2 = _mocked(_api())
        with p1, p2:
            r = self.client.get(URL)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'id="su-file"')
        with mock.patch("explore.views.mint_for", side_effect=FnalLinkRequired("x")):
            self.assertEqual(self.client.get(URL).status_code, 302)
            self.assertEqual(self.client.post(URL, {"step": "file"}).status_code, 401)
        with override_settings(HWDB_WRITE_INSTANCES=[]):
            self.assertEqual(self.client.get(URL).status_code, 403)

    def test_file_makes_a_job_with_the_auto_mapping(self):
        r = self.upload(_api())
        self.assertEqual(r.status_code, 200, r.content)
        job = SheetJob.objects.get()
        self.assertEqual(r.json()["url"], f"{URL}{job.pk}/")
        self.assertEqual((job.name, job.username, job.instance), ("items.csv", "w", "dev"))
        self.assertEqual(job.mapping["S:Vbd"], "spec:Vbd")
        self.assertEqual(job.mapping["Extra"], "")
        p1, p2 = _mocked(_api())
        with p1, p2:
            r = self.client.get(f"{URL}{job.pk}/")
        self.assertContains(r, 'name="col6"')
        self.assertContains(r, "Sheet default: <b>Institution</b> = (186) BNL")
        self.assertContains(r, 'value="pos:A2"')
        self.assertNotContains(r, 'id="su-plan"')

    def test_workbook_tabs_switch_on_the_job_page(self):
        wb = openpyxl.Workbook()
        wb.active.title = "Notes"
        wb.active.append(["just a note"])
        wb.create_sheet("Items").append(["Serial Number", "S:Vbd"])
        wb["Items"].append(["HPK-1", 1])
        wb.create_sheet("Tests").append(["Serial Number", "T:x"])
        wb["Tests"].append(["HPK-1", 1])
        buf = io.BytesIO()
        wb.save(buf)
        api = _api()
        r = self.upload(api, "book.xlsx", buf.getvalue())
        job = SheetJob.objects.get()
        self.assertEqual((job.name, job.sheet["name"], [t["name"] for t in job.tabs]),
                         ("book.xlsx", "Items", ["Notes", "Items", "Tests"]))   # the first tab with rows
        self.assertEqual(job.mapping, {"Serial Number": "serial_number", "S:Vbd": "spec:Vbd"})
        job.rows = [{"n": 2, "state": "done"}]
        job.save()
        url = f"{URL}{job.pk}/"
        p1, p2 = _mocked(api)
        with p1, p2:
            r = self.client.get(url)
            self.assertContains(r, 'id="su-tab"')
            self.assertContains(r, "Tests (1 row)")
            self.assertEqual(self.client.post(url, {"step": "tab", "tab": "Tests"}).status_code, 302)
        job.refresh_from_db()
        self.assertEqual((job.sheet["name"], job.kind, job.mapping, job.rows),
                         ("Tests", "test", {"Serial Number": "serial_number", "T:x": "test:x"}, []))

    def test_bad_files(self):
        self.assertEqual(self.upload(_api(), "a.csv", b"\n\n").status_code, 400)
        p1, p2 = _mocked(_api())
        with p1, p2:
            self.assertEqual(self.client.post(URL, {"step": "file"}).status_code, 400)

    def test_map_plans_apply_writes_and_delete(self):
        api = _api()
        self.upload(api)
        job = SheetJob.objects.get()
        url = f"{URL}{job.pk}/"
        p1, p2 = _mocked(api)
        with p1, p2:
            r = self.client.post(url, {"step": "map", "col0": "part_id", "col1": "serial_number", "col2": "status",
                                       "col3": "spec:Vbd", "col4": "pos:A1", "col5": "comments", "col6": ""})
            self.assertEqual(r.status_code, 302)
            job.refresh_from_db()
            self.assertEqual([(x["action"], x["key"]) for x in job.rows], [("create", "HPK-9"), ("patch", f"{T}-00001")])
            r = self.client.get(url)
            self.assertContains(r, 'id="su-plan"')
            self.assertContains(r, "new item")
            self.assertContains(r, "status → QA/QC Tests - Passed All")
            api.get_subcomponents.return_value = {"data": []}
            api.patch_subcomponents.return_value = _ok()
            api.find_components_by_serial.return_value = []
            r = self.client.post(url, {"step": "apply"})
        self.assertEqual(r.status_code, 200)
        j = r.json()
        self.assertEqual(j["left"], 0)
        self.assertEqual([(x["pid"], x["state"], x["done"]) for x in j["rows"]],
                         [(f"{T}-00009", "done", ["created"]), (f"{T}-00001", "done", ["updated", "positions"])])
        self.assertEqual(api.create_component.call_args.args[1]["serial_number"], "HPK-9")
        self.assertEqual(api.patch_component.call_args.args[1]["status"], {"id": 120})
        job.refresh_from_db()
        self.assertEqual(job.counts()["pending"], 0)
        self.assertIn("1 D00400300001 items created, 1 updated, 0 failed", ActivityEvent.objects.get().summary)
        p1, p2 = _mocked(api)
        with p1, p2:
            r = self.client.get(url)
            self.assertContains(r, "Continue")
            self.assertContains(r, "created")
            self.assertEqual(self.client.post(url, {"step": "delete"}).status_code, 302)
        self.assertFalse(SheetJob.objects.exists())

    def test_apply_records_refusals_and_outages_per_row(self):
        api = _api()
        self.upload(api)
        job = SheetJob.objects.get()
        job.rows = [{"n": 6, "rows": [6], "key": "HPK-9", "pid": "", "action": "create", "changes": [], "error": "",
                     "state": "pending", "rec": {"serial_number": "HPK-9", "institution": {"id": 186, "country_code": "US"},
                                                 "specs": {}, "positions": {}}},
                    {"n": 7, "rows": [7], "key": f"{T}-00001", "pid": f"{T}-00001", "action": "patch", "changes": [],
                     "error": "", "state": "pending", "rec": {"serial_number": "", "status_id": 120, "specs": {}, "positions": {}}}]
        job.save()
        api.find_components_by_serial.return_value = []
        api.create_component.return_value = {"status": "ERROR", "data": "roles"}
        api.get_component.side_effect = requests.ConnectionError("down")
        p1, p2 = _mocked(api)
        with p1, p2:
            j = self.client.post(f"{URL}{job.pk}/", {"step": "apply"}).json()
        self.assertEqual([(x["state"], x["error"]) for x in j["rows"]], [("error", "roles"), ("error", "down")])
        job.refresh_from_db()
        self.assertEqual((job.counts()["error"], job.counts()["failed"]), (0, 2))   # plan problems vs apply failures
        self.assertEqual([x["state"] for x in job.rows], ["error", "error"])

    def test_old_jobs_expire(self):
        from datetime import timedelta
        from django.utils import timezone
        old = SheetJob.objects.create(instance="dev", part_type_id=T, username="w", name="old.csv",
                                      sheet={"values": {}, "columns": ["a"], "rows": []})
        SheetJob.objects.filter(pk=old.pk).update(updated_at=timezone.now() - timedelta(days=4))
        fresh = SheetJob.objects.create(instance="dev", part_type_id=T, username="w", name="fresh.csv",
                                        sheet={"values": {}, "columns": ["a"], "rows": []})
        p1, p2 = _mocked(_api())
        with p1, p2:
            r = self.client.get(URL)
        self.assertContains(r, "fresh.csv")
        self.assertNotContains(r, "old.csv")
        self.assertContains(r, "kept 3 days")
        self.assertEqual(list(SheetJob.objects.values_list("pk", flat=True)), [fresh.pk])

    def test_test_sheet_end_to_end(self):
        api = _api()
        api.get_test_types.return_value = {"data": []}
        api.post_test_type.return_value = _ok()
        api.get_tests.return_value = {"data": []}
        api.post_test.return_value = _ok()
        body = (b"Record Type,Test\r\nTest Name,QC\r\n\r\n"
                b"Serial Number,T:Gain,Comments\r\nHPK-1,12,a\r\nHPK-1,13,b\r\nHPK-9,1,\r\n")
        self.upload(api, "qc.csv", body)
        job = SheetJob.objects.get()
        self.assertEqual((job.kind, job.test_name, job.mapping["T:Gain"]), ("test", "QC", "test:Gain"))
        url = f"{URL}{job.pk}/"
        p1, p2 = _mocked(api)
        with p1, p2:
            r = self.client.get(url)
            self.assertContains(r, 'id="su-kind"')
            self.assertContains(r, 'value="QC"')
            self.assertContains(r, "Test: Gain")
            self.assertNotContains(r, "Specs: Vbd")
            r = self.client.post(url, {"step": "map", "col0": "serial_number", "col1": "test:Gain", "col2": "comments",
                                       "test_name": "QC2"})
            self.assertEqual(r.status_code, 302)
            job.refresh_from_db()
            self.assertEqual([(x["action"], x["pid"], x["rec"].get("test_name")) for x in job.rows],
                             [("test", f"{T}-00001", "QC2"), ("test", f"{T}-00001", "QC2"), ("error", "", "QC2")])
            r = self.client.get(url)
            self.assertContains(r, "test records to post")
            self.assertContains(r, "QC2: Gain = 13")
            # the second post already exists by the time the apply reaches it
            api.get_tests.side_effect = [{"data": []}, {"data": [{"test_data": {"DATA": {"Gain": 13}}}]}]
            api.get_test_types.side_effect = [{"data": []}, {"data": [{"id": 9, "name": "QC2"}]}]
            r = self.client.post(url, {"step": "apply"})
        j = r.json()
        self.assertEqual([(x["state"], x["done"]) for x in j["rows"]], [("done", ["test posted"]), ("done", ["already recorded"])])
        self.assertEqual(api.post_test_type.call_args.args[1]["name"], "QC2")   # created once
        self.assertEqual(api.post_test.call_count, 1)
        self.assertEqual(api.post_test.call_args.args[1]["test_data"], {"DATA": {"Gain": 12}})
        self.assertIn("1 “QC2” test records posted", ActivityEvent.objects.get().summary)
        api.get_test_types.side_effect = None
        p1, p2 = _mocked(api)
        with p1, p2:
            self.assertEqual(self.client.post(url, {"step": "kind", "kind": "item"}).status_code, 302)
        job.refresh_from_db()
        self.assertEqual((job.kind, job.rows, job.mapping["T:Gain"]), ("item", [], "test:Gain"))

    def test_image_sheet_end_to_end(self):
        api = _api()
        api.get_images.return_value = {"data": []}
        api.post_component_image.return_value = _ok()
        body = b"External ID,Image File,Comments\r\nD00400300001-00001,front.png,front\r\nD00400300001-00002,back.png,\r\n"
        self.upload(api, "photos.csv", body)
        job = SheetJob.objects.get()
        self.assertEqual((job.kind, job.mapping), ("image", {"External ID": "part_id", "Image File": "image_file", "Comments": "comments"}))
        url = f"{URL}{job.pk}/"
        p1, p2 = _mocked(api)
        with p1, p2:
            self.assertContains(self.client.get(url), 'value="image" checked')
            self.assertEqual(self.client.post(url, {"step": "map", "col0": "part_id", "col1": "image_file", "col2": "comments"}).status_code, 302)
            r = self.client.get(url)
            self.assertContains(r, 'id="su-files"')
            self.assertContains(r, "<b>2</b> files to attach")
            self.assertContains(r, '"file": "front.png"')
            r = self.client.post(url, {"step": "apply", "n": "3", "last": "1", "file": _file("back.png", b"\x89PNG")})
        j = r.json()
        self.assertEqual([(x["n"], x["state"], x["done"]) for x in j["rows"]], [(3, "done", ["attached"])])
        self.assertEqual(j["left"], 1)
        a = api.post_component_image.call_args
        self.assertEqual((a.args[0], a.args[2], a.args[3], a.args[4]), (f"{T}-00002", "back.png", "Sheet upload via HWDB Explorer", "image/png"))
        self.assertIn("1 files attached", ActivityEvent.objects.get().summary)
        p1, p2 = _mocked(api)
        with p1, p2:
            self.assertEqual(self.client.post(url, {"step": "apply", "n": "2"}).status_code, 400)   # no file

    def test_other_users_jobs_are_invisible(self):
        SheetJob.objects.create(instance="dev", part_type_id=T, username="someone", name="x.csv",
                                sheet={"values": {}, "columns": ["a"], "rows": []})
        p1, p2 = _mocked(_api())
        with p1, p2:
            self.assertNotContains(self.client.get(URL), "x.csv")
            self.assertEqual(self.client.get(f"{URL}{SheetJob.objects.get().pk}/").status_code, 302)


def _file(name, body):
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile(name, body)
